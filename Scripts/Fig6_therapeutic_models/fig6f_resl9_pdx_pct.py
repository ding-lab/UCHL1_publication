"""
Figure 6f — RESL9 PDX (BAP1-WT), relative tumor growth (%)

Reads:   RESL9C_PDX_relative_growth_data_V2.xlsx
         sheets: Summary_relative_growth, Per_mouse_relative_growth, Two_way_ANOVA
Writes:  fig6f_resl9_pdx_pct.png and .pdf
"""

import warnings; warnings.filterwarnings('ignore')
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import openpyxl
import statsmodels.api as sm
from statsmodels.formula.api import ols

# --------- Configuration ---------
DATA_XLSX = './ccRCC_UCHL1_manuscript/Data/RESL9C_PDX_relative_growth_data_V2.xlsx'
OUT_PNG = './ccRCC_UCHL1_manuscript/Data/fig6f_resl9_pdx_pct.png'
OUT_PDF = './ccRCC_UCHL1_manuscript/Data/fig6f_resl9_pdx_pct.pdf'

REF_GROUP = 'Vehicle'
GROUP_ORDER = ['Vehicle', 'IMP-1710', 'Cabo', 'Cabo+IMP-1710']
STYLE = {
    'Vehicle':       {'color': '#1f77b4', 'marker': 'o'},
    'IMP-1710':      {'color': '#e8000b', 'marker': 's'},
    'Cabo':          {'color': '#2ca02c', 'marker': '^'},
    'Cabo+IMP-1710': {'color': '#7e2bd4', 'marker': 'v'},
}
TITLE = 'RESL9 PDX (BAP1-WT) relative tumor growth'
XLABEL = 'Time (days after treatment start)'
YLABEL = 'Relative tumor growth (%)'
NORMALIZATION_NOTE = "Each mouse normalized to its own day-0 volume (baseline = 100%)."


def tier(p):
    if p is None or (isinstance(p, float) and np.isnan(p)): return 'na'
    if p < 0.0001: return '****'
    if p < 0.001:  return '***'
    if p < 0.01:   return '**'
    if p < 0.05:   return '*'
    return 'ns'


def read_table(xlsx, sheet, header_row=3):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[sheet]
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {}
        for c_idx, h in enumerate(headers, start=1):
            if h is None: continue
            row[h] = ws.cell(r, c_idx).value
        if any(v is not None for v in row.values()):
            rows.append(row)
    return pd.DataFrame(rows)


def two_way(df, ref, other, value_col):
    ref_days = set(df[df['Group'] == ref]['Day'].unique())
    oth_days = set(df[df['Group'] == other]['Day'].unique())
    shared = sorted(ref_days & oth_days)
    sub = df[df['Group'].isin([ref, other]) & df['Day'].isin(shared)].copy()
    sub['Group'] = sub['Group'].astype(str)
    sub = sub.rename(columns={value_col: 'value'})
    model = ols('value ~ C(Group) * C(Day)', data=sub).fit()
    a = sm.stats.anova_lm(model, typ=2)
    return {
        'Comparison': f'{other} vs {ref}',
        'Interaction F': float(a.loc['C(Group):C(Day)', 'F']),
        'Interaction p': float(a.loc['C(Group):C(Day)', 'PR(>F)']),
        'Interaction df': int(a.loc['C(Group):C(Day)', 'df']),
        'Residual df': int(a.loc['Residual', 'df']),
        'Days used': shared,
    }

print(f'Reading {DATA_XLSX}...')
summary = read_table(DATA_XLSX, 'Summary_relative_growth')
permouse = read_table(DATA_XLSX, 'Per_mouse_relative_growth')
stored_anova = read_table(DATA_XLSX, 'Two_way_ANOVA')

if 'Cell line / Group' in summary.columns:
    summary = summary.rename(columns={'Cell line / Group': 'Group'})

for col in ['Day', 'Mean relative tumor growth', 'SD across mice', 'SEM across mice', 'n_mice']:
    if col in summary.columns:
        summary[col] = pd.to_numeric(summary[col], errors='coerce')
for col in ['Day', 'Relative tumor growth', 'Mean of bilateral flanks (mm³)', 'Baseline volume (mm³)']:
    if col in permouse.columns:
        permouse[col] = pd.to_numeric(permouse[col], errors='coerce')

summary = summary.dropna(subset=['Day']).copy()
summary['Day'] = summary['Day'].astype(int)
permouse = permouse.dropna(subset=['Day', 'Relative tumor growth']).copy()
permouse['Day'] = permouse['Day'].astype(int)

summary['mean_pct'] = summary['Mean relative tumor growth'] * 100
summary['sem_pct']  = summary['SEM across mice']           * 100

computed_p = {}
for other in [g for g in GROUP_ORDER if g != REF_GROUP]:
    res = two_way(permouse, REF_GROUP, other, 'Relative tumor growth')
    computed_p[other] = res['Interaction p']

plt.rcParams.update({
    'font.family': 'DejaVu Sans', 'font.size': 11,
    'axes.labelsize': 13, 'axes.titlesize': 13, 'axes.titleweight': 'bold',
    'xtick.labelsize': 11, 'ytick.labelsize': 11, 'legend.fontsize': 10.5,
    'axes.spines.top': False, 'axes.spines.right': False,
})
fig, ax = plt.subplots(figsize=(8.0, 5.5))

days_all = sorted(summary['Day'].unique())
for g in GROUP_ORDER:
    s = summary[summary['Group'] == g].sort_values('Day')
    if s.empty: continue
    st = STYLE[g]
    ax.errorbar(s['Day'].values, s['mean_pct'].values,
                yerr=s['sem_pct'].values,
                color=st['color'], marker=st['marker'], linestyle='-',
                linewidth=1.8, markersize=7,
                markeredgecolor='white', markeredgewidth=0.8,
                capsize=4, capthick=1.1, elinewidth=1.1, label=g, zorder=3)

ax.set_xlabel(XLABEL); ax.set_ylabel(YLABEL); ax.set_xticks(days_all)
ax.set_title(TITLE, pad=8)
ax.yaxis.grid(True, linestyle='--', alpha=0.35)

final_day = days_all[-1]
ref_row = summary[(summary['Group'] == REF_GROUP) & (summary['Day'] == final_day)]
if len(ref_row):
    ref_top = ref_row.iloc[0]['mean_pct'] + (ref_row.iloc[0]['sem_pct'] or 0)
    drawn = [g for g in GROUP_ORDER if g != REF_GROUP and computed_p.get(g, np.nan) < 0.05]
    all_tops = [ref_top]
    for g in GROUP_ORDER:
        s = summary[summary['Group'] == g]
        all_tops.append((s['mean_pct'] + s['sem_pct'].fillna(0)).max())
    ax.set_ylim(0, max(all_tops) * (1.28 if drawn else 1.12))
    span = final_day - days_all[0]
    if drawn:
        tip = span * 0.022; step = span * 0.085
        xs = [final_day + span * 0.07 + i * step for i in range(len(drawn))]
        for x_b, g in zip(xs, drawn):
            row = summary[(summary['Group'] == g) & (summary['Day'] == final_day)].iloc[0]
            g_top = row['mean_pct'] + (row['sem_pct'] or 0)
            ax.plot([x_b - tip, x_b, x_b, x_b - tip],
                    [g_top, g_top, ref_top, ref_top], color='black', lw=1.2, zorder=4)
            ax.text(x_b + span * 0.02, (g_top + ref_top) / 2, tier(computed_p[g]),
                    ha='left', va='center', fontsize=10, fontweight='bold', zorder=5)
        ax.set_xlim(days_all[0] - span * 0.05, xs[-1] + span * 0.14)
    else:
        ax.set_xlim(days_all[0] - span * 0.05, final_day + span * 0.06)

ax.legend(loc='upper left', frameon=False, fontsize=10.5)

n_per_group = {g: int(summary[summary['Group']==g]['n_mice'].iloc[0])
               for g in GROUP_ORDER if not summary[summary['Group']==g].empty}
n_text = '; '.join(f"{g} n={n_per_group[g]}" for g in GROUP_ORDER if g in n_per_group)
fig.text(0.04, 0.005,
         'Two-way ANOVA interaction term (group × time, per-mouse relative growth); '
         '*p<0.05  **p<0.01  ***p<0.001  ****p<0.0001',
         fontsize=8, color='#555555')
fig.text(0.04, -0.018, f"{n_text}. Error bars = SEM across mice. {NORMALIZATION_NOTE}",
         fontsize=8, color='#555555')
plt.tight_layout()
plt.savefig(OUT_PNG, dpi=300, bbox_inches='tight')
plt.savefig(OUT_PDF, bbox_inches='tight')
plt.close(fig)